from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

from app.services import document_loader
from app.services.document_loader import iter_office_embedded_images, load_raw_documents


def test_image_ocr_loader_extracts_text_and_metadata(
    tmp_path: Path,
    monkeypatch,
) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = raw_data_dir / "screenshot.png"
    Image.new("RGB", (320, 120), color="white").save(image_path)

    monkeypatch.setattr(document_loader, "_tesseract_available", lambda: True)
    monkeypatch.setattr(
        document_loader.pytesseract,
        "image_to_string",
        lambda image, lang: "Найти термин и начать обучение",
    )

    documents = load_raw_documents(raw_data_dir, tenant_id="tenant", bucket_id="bucket")

    image_documents = [document for document in documents if document.source_type == "image_ocr"]

    assert len(image_documents) == 1
    assert image_documents[0].tenant_id == "tenant"
    assert image_documents[0].bucket_id == "bucket"
    assert "Найти термин" in image_documents[0].content
    assert image_documents[0].metadata["image_width"] == 320
    assert image_documents[0].metadata["image_height"] == 120
    assert image_documents[0].metadata["ocr_engine"] == "tesseract"
    assert image_documents[0].metadata["ocr_languages"] == "rus+eng"


def test_docx_embedded_image_ocr_builds_parent_metadata(
    tmp_path: Path,
    monkeypatch,
) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = tmp_path / "embedded.png"
    Image.new("RGB", (100, 80), color="white").save(image_path)

    docx_path = raw_data_dir / "with-image.docx"
    document = Document()
    document.add_paragraph("Document with embedded image.")
    document.add_picture(str(image_path))
    document.save(docx_path)

    monkeypatch.setattr(document_loader, "_tesseract_available", lambda: True)
    monkeypatch.setattr(
        document_loader.pytesseract,
        "image_to_string",
        lambda image, lang: "Embedded chart text",
    )

    documents = load_raw_documents(raw_data_dir, tenant_id="tenant", bucket_id="bucket")
    embedded_ocr_documents = [
        document
        for document in documents
        if document.source_type == "image_ocr"
        and document.metadata.get("parent_source_type") == "docx"
    ]

    assert len(embedded_ocr_documents) == 1
    assert embedded_ocr_documents[0].source_path == str(docx_path)
    assert "Embedded chart text" in embedded_ocr_documents[0].content
    assert embedded_ocr_documents[0].metadata["embedded_path"].startswith("word/media/")
    assert embedded_ocr_documents[0].metadata["embedded_image_index"] == 1


def test_docx_table_cell_image_extractor_builds_anchor_metadata(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = tmp_path / "embedded.png"
    Image.new("RGB", (100, 80), color="white").save(image_path)

    docx_path = raw_data_dir / "with-table-image.docx"
    document = Document()
    table = document.add_table(rows=2, cols=4)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Price"
    table.rows[0].cells[2].text = "Params"
    table.rows[0].cells[3].text = "Image"
    table.rows[1].cells[0].text = "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ"
    table.rows[1].cells[1].text = "Цена за кегу 3500 р."
    table.rows[1].cells[2].text = "4,1 % алкоголь, 11% плотность"
    table.rows[1].cells[3].paragraphs[0].add_run().add_picture(str(image_path))
    document.save(docx_path)

    embedded_images = iter_office_embedded_images(raw_data_dir)

    assert len(embedded_images) == 1
    assert embedded_images[0].metadata["anchor_type"] == "docx_table_cell"
    assert embedded_images[0].metadata["table_index"] == 0
    assert embedded_images[0].metadata["table_row_index"] == 1
    assert embedded_images[0].metadata["table_cell_index"] == 3
    assert "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ" in embedded_images[0].metadata["linked_text"]
    assert "3500" in embedded_images[0].metadata["linked_text"]


def test_xlsx_image_extractor_builds_anchor_metadata(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = tmp_path / "embedded.png"
    Image.new("RGB", (80, 60), color="white").save(image_path)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Прайс"
    worksheet["B2"] = "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ"
    worksheet["C2"] = "Цена за кегу 3500 р."
    worksheet["D2"] = "4,1 % алкоголь, 11% плотность"
    worksheet.add_image(ExcelImage(str(image_path)), "B2")
    xlsx_path = raw_data_dir / "with-anchor.xlsx"
    workbook.save(xlsx_path)

    embedded_images = iter_office_embedded_images(raw_data_dir)

    assert len(embedded_images) == 1
    assert embedded_images[0].parent_path == xlsx_path
    assert embedded_images[0].parent_source_type == "xlsx"
    assert embedded_images[0].metadata["anchor_type"] == "xlsx_cell"
    assert embedded_images[0].metadata["sheet_name"] == "Прайс"
    assert embedded_images[0].metadata["anchor_row"] == 2
    assert embedded_images[0].metadata["anchor_col"] == 2
    assert embedded_images[0].metadata["anchor_cell"] == "B2"
    assert "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ" in embedded_images[0].metadata["linked_text"]
    assert "3500" in embedded_images[0].metadata["linked_text"]


def test_legacy_xls_embedded_image_extractor_finds_binary_image_blobs(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_buffer = tmp_path / "image.jpg"
    Image.new("RGB", (64, 48), color="white").save(image_buffer, format="JPEG")

    xls_path = raw_data_dir / "legacy.xls"
    xls_path.write_bytes(b"legacy-biff-prefix" + image_buffer.read_bytes() + b"legacy-biff-suffix")

    embedded_images = iter_office_embedded_images(raw_data_dir)

    assert len(embedded_images) == 1
    assert embedded_images[0].parent_path == xls_path
    assert embedded_images[0].parent_source_type == "xls"
    assert embedded_images[0].embedded_path == "legacy-binary/image1.jpg"
    assert embedded_images[0].image_index == 1
