from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image, ImageDraw

from app.services import document_loader
from app.services.image_digest_service import load_image_digest_documents


class _FakeOllamaClient:
    def __init__(self) -> None:
        self.generate_kwargs = {}

    async def generate(self, **kwargs):
        self.generate_kwargs = kwargs
        return {"response": "На изображении видна надпись UFA и человек рядом."}


def test_load_image_digest_documents_builds_digest_metadata(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = raw_data_dir / "photo.png"
    Image.new("RGB", (640, 480), color="white").save(image_path)
    ollama_client = _FakeOllamaClient()

    async def run_loader():
        return await load_image_digest_documents(
            raw_data_dir,
            ollama_client=ollama_client,
            vision_model="vision-model",
            tenant_id="tenant",
            bucket_id="bucket",
        )

    import asyncio

    documents = asyncio.run(run_loader())

    assert len(documents) == 1
    assert documents[0].source_type == "image_digest"
    assert documents[0].tenant_id == "tenant"
    assert documents[0].bucket_id == "bucket"
    assert "надпись UFA" in documents[0].content
    assert documents[0].metadata["image_width"] == 640
    assert documents[0].metadata["image_height"] == 480
    assert documents[0].metadata["vision_model"] == "vision-model"
    assert documents[0].metadata["digest_type"] == "vision_caption"
    assert ollama_client.generate_kwargs["images"]


def test_load_image_digest_documents_includes_docx_embedded_images(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = tmp_path / "embedded.png"
    Image.new("RGB", (200, 100), color="white").save(image_path)

    docx_path = raw_data_dir / "with-image.docx"
    document = Document()
    document.add_paragraph("Document with embedded image.")
    document.add_picture(str(image_path))
    document.save(docx_path)
    ollama_client = _FakeOllamaClient()

    async def run_loader():
        return await load_image_digest_documents(
            raw_data_dir,
            ollama_client=ollama_client,
            vision_model="vision-model",
            tenant_id="tenant",
            bucket_id="bucket",
        )

    import asyncio

    documents = asyncio.run(run_loader())

    assert len(documents) == 1
    assert documents[0].source_type == "image_digest"
    assert documents[0].source_path == str(docx_path)
    assert documents[0].metadata["parent_source_type"] == "docx"
    assert documents[0].metadata["embedded_path"].startswith("word/media/")
    assert documents[0].metadata["embedded_image_index"] == 1
    assert documents[0].metadata["anchor_type"] == "docx_paragraph"
    assert documents[0].metadata["previous_paragraph_text"] == "Document with embedded image."
    assert documents[0].metadata["linked_text"] == "Document with embedded image."
    assert "Linked text: Document with embedded image." in documents[0].content
    assert documents[0].metadata["image_width"] == 200
    assert documents[0].metadata["image_height"] == 100
    assert ollama_client.generate_kwargs["images"]


def test_load_image_digest_documents_includes_legacy_xls_embedded_images(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = tmp_path / "embedded.jpg"
    Image.new("RGB", (120, 90), color="white").save(image_path, format="JPEG")
    xls_path = raw_data_dir / "legacy.xls"
    xls_path.write_bytes(b"legacy-biff-prefix" + image_path.read_bytes() + b"legacy-biff-suffix")
    ollama_client = _FakeOllamaClient()

    async def run_loader():
        return await load_image_digest_documents(
            raw_data_dir,
            ollama_client=ollama_client,
            vision_model="vision-model",
            tenant_id="tenant",
            bucket_id="bucket",
        )

    import asyncio

    documents = asyncio.run(run_loader())

    assert len(documents) == 1
    assert documents[0].source_type == "image_digest"
    assert documents[0].source_path == str(xls_path)
    assert documents[0].metadata["parent_source_type"] == "xls"
    assert documents[0].metadata["embedded_path"] == "legacy-binary/image1.jpg"
    assert documents[0].metadata["embedded_image_index"] == 1
    assert documents[0].metadata["image_width"] == 120
    assert documents[0].metadata["image_height"] == 90
    assert ollama_client.generate_kwargs["images"]


def test_load_image_digest_documents_includes_xlsx_anchor_context(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = tmp_path / "embedded.png"
    Image.new("RGB", (80, 60), color="white").save(image_path)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Прайс"
    worksheet["B2"] = "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ"
    worksheet["C2"] = "Цена за кегу 3500 р."
    worksheet["D2"] = "4,1 % алкоголь, 11% плотность"
    worksheet.add_image(ExcelImage(str(image_path)), "B2")
    xlsx_path = raw_data_dir / "with-anchor.xlsx"
    workbook.save(xlsx_path)
    ollama_client = _FakeOllamaClient()

    async def run_loader():
        return await load_image_digest_documents(
            raw_data_dir,
            ollama_client=ollama_client,
            vision_model="vision-model",
            tenant_id="tenant",
            bucket_id="bucket",
        )

    import asyncio

    documents = asyncio.run(run_loader())

    assert len(documents) == 1
    assert documents[0].source_path == str(xlsx_path)
    assert documents[0].metadata["parent_source_type"] == "xlsx"
    assert documents[0].metadata["anchor_type"] == "xlsx_cell"
    assert documents[0].metadata["sheet_name"] == "Прайс"
    assert documents[0].metadata["anchor_cell"] == "B2"
    assert "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ" in documents[0].metadata["linked_text"]
    assert "Linked text: ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ" in documents[0].content
    assert ollama_client.generate_kwargs["images"]


def test_load_image_digest_documents_includes_scanned_table_anchor_context(
    tmp_path: Path,
    monkeypatch,
) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    image_path = raw_data_dir / "scanned_table.png"
    image = Image.new("RGB", (500, 260), color="white")
    draw = ImageDraw.Draw(image)
    grid_color = (60, 120, 210)
    columns = [20, 160, 300, 420, 480]
    rows = [20, 60, 140, 220]
    for x in columns:
        draw.line((x, rows[0], x, rows[-1]), fill=grid_color, width=2)
    for y in rows:
        draw.line((columns[0], y, columns[-1], y), fill=grid_color, width=2)
    draw.text((30, 75), "PYA TNITSKOE", fill="black")
    draw.text((170, 75), "3500 r", fill="black")
    draw.text((310, 75), "4.1 11", fill="black")
    draw.rectangle((430, 80, 470, 125), fill=(160, 40, 40), outline=(220, 180, 60), width=3)
    image.save(image_path)
    monkeypatch.setattr(
        document_loader,
        "_ocr_image_region",
        lambda image: "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ Цена за кегу 3500 р. 4,1 % алкоголь, 11% плотность",
    )
    ollama_client = _FakeOllamaClient()

    async def run_loader():
        return await load_image_digest_documents(
            raw_data_dir,
            ollama_client=ollama_client,
            vision_model="vision-model",
            tenant_id="tenant",
            bucket_id="bucket",
        )

    import asyncio

    documents = asyncio.run(run_loader())
    scanned_documents = [
        document
        for document in documents
        if document.metadata.get("parent_source_type") == "scanned_table"
    ]

    assert len(scanned_documents) == 1
    assert scanned_documents[0].source_path == str(image_path)
    assert scanned_documents[0].metadata["anchor_type"] == "scanned_table_cell"
    assert "ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ" in scanned_documents[0].metadata["linked_text"]
    assert "Linked text: ПЯТНИЦКОЕ НЕФИЛЬТРОВАННОЕ" in scanned_documents[0].content
    assert ollama_client.generate_kwargs["images"]
