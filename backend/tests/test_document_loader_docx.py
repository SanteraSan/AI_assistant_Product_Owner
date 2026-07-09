from pathlib import Path

from docx import Document

from app.services.document_loader import load_raw_documents


def test_docx_loader_extracts_paragraphs_and_table_rows(tmp_path: Path) -> None:
    raw_data_dir = tmp_path / "raw"
    raw_data_dir.mkdir()
    docx_path = raw_data_dir / "brief.docx"

    document = Document()
    document.add_paragraph("Enterprise onboarding risk.")
    document.add_paragraph("Task 9. Improve ModifyUsers performance.")
    document.add_paragraph("async Task ModifyUsers(int userId, int[] usersIds, UserStatus status)")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Area"
    table.rows[0].cells[1].text = "Owner"
    row = table.add_row().cells
    row[0].text = "DOCX ingestion"
    row[1].text = "AI Platform"
    document.save(docx_path)

    documents = load_raw_documents(raw_data_dir, tenant_id="tenant", bucket_id="bucket")

    docx_documents = [document for document in documents if document.source_type == "docx"]
    block_types = {document.metadata["block_type"] for document in docx_documents}

    assert len(docx_documents) == 6
    assert block_types == {"paragraph", "paragraph_window", "table_row"}
    assert {document.tenant_id for document in docx_documents} == {"tenant"}
    assert {document.bucket_id for document in docx_documents} == {"bucket"}
    assert any("Enterprise onboarding risk" in document.content for document in docx_documents)
    assert any("DOCX ingestion | AI Platform" in document.content for document in docx_documents)
    assert any(
        document.metadata["block_type"] == "paragraph_window"
        and "Task 9. Improve ModifyUsers performance" in document.content
        and "async Task ModifyUsers" in document.content
        for document in docx_documents
    )
