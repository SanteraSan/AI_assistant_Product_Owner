from __future__ import annotations

import warnings
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path
from typing import Any, Iterator
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from app.services.loaders import common
from app.services.loaders.common import (
    DOCX_MEDIA_PREFIX,
    IMAGE_OCR_LANGUAGES,
    LEGACY_XLS_EMBEDDED_IMAGE_LIMIT,
    OFFICE_IMAGE_EXTENSIONS,
    OFFICE_XML_NAMESPACES,
    XLSX_MEDIA_PREFIX,
    EmbeddedImage,
    RawDocument,
    _domain_for_path,
    _embedded_image_context_lines,
    _features_from_text,
    _open_embedded_image,
    _stable_document_id,
)


def _load_office_embedded_image_ocr_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    if not common._tesseract_available():
        return []

    documents: list[RawDocument] = []
    for embedded_image in iter_office_embedded_images(raw_data_dir):
        try:
            image, width, height, image_format = _open_embedded_image(embedded_image)
            content = common.ocr_image_text(image)
        except Exception as exc:
            warnings.warn(
                f"Embedded image OCR skipped for {embedded_image.parent_path}:{embedded_image.embedded_path}: {exc}",
                RuntimeWarning,
                stacklevel=2,
            )
            continue
        if not content:
            continue

        parent_path = embedded_image.parent_path
        documents.append(
            RawDocument(
                id=f"{_stable_document_id(parent_path)}:embedded_image:{embedded_image.image_index}:ocr",
                title=f"{parent_path.stem} - embedded image {embedded_image.image_index} OCR text",
                content="\n".join(
                    [
                        f"File: {parent_path.name}",
                        f"Embedded image: {embedded_image.embedded_path}",
                        "Block type: image_ocr",
                        *_embedded_image_context_lines(embedded_image.metadata),
                        "OCR text:",
                        content,
                    ]
                ),
                source_type="image_ocr",
                source_path=str(parent_path),
                domain=_domain_for_path(parent_path),
                feature=_features_from_text(content),
                metadata={
                    "file_name": parent_path.name,
                    "block_type": "image_ocr",
                    "parent_source_type": embedded_image.parent_source_type,
                    "embedded_path": embedded_image.embedded_path,
                    "embedded_image_index": embedded_image.image_index,
                    "image_width": width,
                    "image_height": height,
                    "image_format": image_format,
                    "ocr_engine": "tesseract",
                    "ocr_languages": IMAGE_OCR_LANGUAGES,
                    **embedded_image.metadata,
                },
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def iter_office_embedded_images(raw_data_dir: Path) -> list[EmbeddedImage]:
    images: list[EmbeddedImage] = []
    for path in sorted(
        list(raw_data_dir.rglob("*.docx"))
        + list(raw_data_dir.rglob("*.xlsx"))
        + list(raw_data_dir.rglob("*.xls"))
    ):
        images.extend(_embedded_images_from_office_file(path))
    return images


def _embedded_images_from_office_file(path: Path) -> list[EmbeddedImage]:
    if path.suffix.lower() == ".xls":
        return _embedded_images_from_legacy_xls(path)
    if path.suffix.lower() == ".xlsx":
        return _embedded_images_from_xlsx(path)
    if path.suffix.lower() == ".docx":
        return _embedded_images_from_docx(path)

    return []


def _embedded_images_from_docx(path: Path) -> list[EmbeddedImage]:
    try:
        with ZipFile(path) as archive:
            metadata_by_path = _docx_image_metadata_by_embedded_path(archive)
            media_paths = sorted(
                name
                for name in archive.namelist()
                if name.startswith(DOCX_MEDIA_PREFIX)
                and Path(name).suffix.lower() in OFFICE_IMAGE_EXTENSIONS
            )
            return [
                EmbeddedImage(
                    parent_path=path,
                    embedded_path=media_path,
                    image_index=image_index,
                    content=archive.read(media_path),
                    parent_source_type="docx",
                    metadata=metadata_by_path.get(media_path, {}),
                )
                for image_index, media_path in enumerate(media_paths, start=1)
            ]
    except BadZipFile:
        warnings.warn(
            f"Office embedded images skipped for non-zip Office file: {path}",
            RuntimeWarning,
            stacklevel=2,
        )
        return []


def _embedded_images_from_xlsx(path: Path) -> list[EmbeddedImage]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    images: list[EmbeddedImage] = []
    image_index = 0
    for worksheet in workbook.worksheets:
        for image in getattr(worksheet, "_images", []):
            image_index += 1
            content = image._data()
            extension = _image_extension(getattr(image, "format", None), content)
            images.append(
                EmbeddedImage(
                    parent_path=path,
                    embedded_path=f"{XLSX_MEDIA_PREFIX}anchored_image{image_index}.{extension}",
                    image_index=image_index,
                    content=content,
                    parent_source_type="xlsx",
                    metadata=_xlsx_image_anchor_metadata(worksheet, image),
                )
            )
    return images


def _embedded_images_from_legacy_xls(path: Path) -> list[EmbeddedImage]:
    data = path.read_bytes()
    image_blobs = [
        *list(_iter_jpeg_blobs(data)),
        *list(_iter_png_blobs(data)),
    ]
    images: list[EmbeddedImage] = []
    for image_index, (extension, content) in enumerate(
        image_blobs[:LEGACY_XLS_EMBEDDED_IMAGE_LIMIT],
        start=1,
    ):
        images.append(
            EmbeddedImage(
                parent_path=path,
                embedded_path=f"legacy-binary/image{image_index}.{extension}",
                image_index=image_index,
                content=content,
                parent_source_type="xls",
            )
        )
    return images


def _docx_image_metadata_by_embedded_path(archive: ZipFile) -> dict[str, dict[str, Any]]:
    try:
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        relationships = _docx_relationship_targets(archive)
    except Exception:
        return {}

    metadata_by_path: dict[str, dict[str, Any]] = {}
    body = document_xml.find("w:body", OFFICE_XML_NAMESPACES)
    if body is None:
        return metadata_by_path

    paragraph_index = -1
    table_index = -1
    previous_paragraph_text = ""
    for child in list(body):
        if child.tag == f"{{{OFFICE_XML_NAMESPACES['w']}}}p":
            paragraph_index += 1
            paragraph_text = _xml_text(child)
            linked_text = paragraph_text or previous_paragraph_text
            for relationship_id in _embedded_relationship_ids(child):
                embedded_path = _docx_embedded_path(relationships.get(relationship_id, ""))
                if embedded_path:
                    metadata_by_path.setdefault(
                        embedded_path,
                        {
                            "anchor_type": "docx_paragraph",
                            "paragraph_index": paragraph_index,
                            "paragraph_text": paragraph_text,
                            "previous_paragraph_text": previous_paragraph_text,
                            "linked_text": linked_text,
                        },
                    )
            if paragraph_text:
                previous_paragraph_text = paragraph_text
            continue

        if child.tag != f"{{{OFFICE_XML_NAMESPACES['w']}}}tbl":
            continue

        table_index += 1
        for row_index, row in enumerate(child.findall("w:tr", OFFICE_XML_NAMESPACES)):
            cells = row.findall("w:tc", OFFICE_XML_NAMESPACES)
            row_values = [_xml_text(cell) for cell in cells]
            row_text = " | ".join(value for value in row_values if value)
            for cell_index, cell in enumerate(cells):
                cell_text = row_values[cell_index] if cell_index < len(row_values) else ""
                for relationship_id in _embedded_relationship_ids(cell):
                    embedded_path = _docx_embedded_path(relationships.get(relationship_id, ""))
                    if embedded_path:
                        metadata_by_path.setdefault(
                            embedded_path,
                            {
                                "anchor_type": "docx_table_cell",
                                "table_index": table_index,
                                "table_row_index": row_index,
                                "table_cell_index": cell_index,
                                "table_row_text": row_text,
                                "table_cell_text": cell_text,
                                "linked_text": row_text or cell_text,
                            },
                        )
    return metadata_by_path


def _docx_relationship_targets(archive: ZipFile) -> dict[str, str]:
    relationships_xml = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    return {
        relationship.attrib["Id"]: relationship.attrib.get("Target", "")
        for relationship in relationships_xml.findall(
            "rel:Relationship",
            OFFICE_XML_NAMESPACES,
        )
        if relationship.attrib.get("Id")
    }


def _embedded_relationship_ids(node: ET.Element) -> list[str]:
    relationship_ids: list[str] = []
    for blip in node.findall(".//a:blip", OFFICE_XML_NAMESPACES):
        relationship_id = blip.attrib.get(f"{{{OFFICE_XML_NAMESPACES['r']}}}embed")
        if relationship_id:
            relationship_ids.append(relationship_id)
    return relationship_ids


def _docx_embedded_path(target: str) -> str:
    if not target:
        return ""
    if target.startswith("word/"):
        return target
    return f"word/{target.lstrip('/')}"


def _xml_text(node: ET.Element) -> str:
    return "".join(
        text.text or ""
        for text in node.findall(".//w:t", OFFICE_XML_NAMESPACES)
    ).strip()


def _xlsx_image_anchor_metadata(worksheet: Any, image: Any) -> dict[str, Any]:
    marker = getattr(getattr(image, "anchor", None), "_from", None)
    if marker is None:
        return {
            "anchor_type": "xlsx_unknown",
            "sheet_name": worksheet.title,
        }

    row_number = int(marker.row) + 1
    column_number = int(marker.col) + 1
    row_text = _worksheet_row_text_by_number(worksheet, row_number)
    cell_value = worksheet.cell(row=row_number, column=column_number).value
    return {
        "anchor_type": "xlsx_cell",
        "sheet_name": worksheet.title,
        "anchor_row": row_number,
        "anchor_col": column_number,
        "anchor_cell": f"{get_column_letter(column_number)}{row_number}",
        "anchor_cell_value": str(cell_value).strip() if cell_value is not None else "",
        "nearby_row_text": row_text,
        "linked_text": row_text,
    }


def _worksheet_row_text_by_number(worksheet: Any, row_number: int) -> str:
    values = [
        str(cell.value).strip()
        for cell in worksheet[row_number]
        if cell.value not in (None, "")
    ]
    return " | ".join(values)


def _image_extension(image_format: str | None, content: bytes) -> str:
    normalized_format = (image_format or "").lower()
    if normalized_format in {"jpeg", "jpg"}:
        return "jpg"
    if normalized_format == "png":
        return "png"
    if content.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    return normalized_format or "bin"


def _iter_jpeg_blobs(data: bytes) -> Iterator[tuple[str, bytes]]:
    position = 0
    while True:
        start = data.find(b"\xff\xd8\xff", position)
        if start < 0:
            return
        end = data.find(b"\xff\xd9", start)
        if end < 0:
            return
        blob = data[start : end + 2]
        if _valid_image_blob(blob):
            yield "jpg", blob
        position = end + 2


def _iter_png_blobs(data: bytes) -> Iterator[tuple[str, bytes]]:
    signature = b"\x89PNG\r\n\x1a\n"
    position = 0
    while True:
        start = data.find(signature, position)
        if start < 0:
            return
        end = data.find(b"IEND", start)
        if end < 0:
            return
        blob = data[start : end + 8]
        if _valid_image_blob(blob):
            yield "png", blob
        position = end + 8


def _valid_image_blob(blob: bytes) -> bool:
    try:
        with Image.open(BytesIO(blob)) as image:
            image.load()
    except Exception:
        return False
    return True

