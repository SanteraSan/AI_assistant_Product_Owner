from __future__ import annotations

import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
import pandas as pd

from app.services.loaders.common import (
    RawDocument,
    _compact_row_values,
    _domain_for_path,
    _features_from_text,
    _row_to_text,
    _stable_document_id,
)


def _load_excel_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    excel_paths = sorted(
        list(raw_data_dir.rglob("*.xlsx"))
        + list(raw_data_dir.rglob("*.xls"))
    )
    for path in excel_paths:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str)
        for sheet_name, frame in sheets.items():
            for index, row in frame.fillna("").iterrows():
                row_data = {key: str(value) for key, value in row.to_dict().items()}
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Sheet: {sheet_name}",
                        f"Row values: {_compact_row_values(row_data)}",
                        _row_to_text(row_data),
                    ]
                )
                title = f"{path.stem} - {sheet_name} - row {int(index) + 2}"
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:{sheet_name}:{index}",
                        title=title,
                        content=content,
                        source_type="excel_row",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "sheet_name": sheet_name,
                            "row_index": int(index),
                            "excel_row_number": int(index) + 2,
                            **row_data,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    return documents


def _load_excel_chart_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.xlsx")):
        workbook = load_workbook(path, data_only=True, read_only=False)
        path_documents: list[RawDocument] = []
        for worksheet in workbook.worksheets:
            charts = getattr(worksheet, "_charts", [])
            if not charts:
                continue
            sheet_context = _worksheet_context(worksheet)
            for chart_index, chart in enumerate(charts, start=1):
                chart_type = type(chart).__name__
                chart_refs = _chart_references(chart)
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Sheet: {worksheet.title}",
                        "Block type: excel_chart",
                        f"Chart type: {chart_type}",
                        f"Chart references: {', '.join(chart_refs) if chart_refs else 'not available'}",
                        "Visible sheet context:",
                        sheet_context,
                    ]
                )
                path_documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:{worksheet.title}:chart:{chart_index}",
                        title=f"{path.stem} - {worksheet.title} - chart {chart_index}",
                        content=content,
                        source_type="excel_chart",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "sheet_name": worksheet.title,
                            "block_type": "excel_chart",
                            "chart_index": chart_index,
                            "chart_type": chart_type,
                            "chart_references": chart_refs,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
        documents.extend(path_documents)
        if not path_documents:
            documents.extend(
                _load_excel_chart_xml_documents(
                    path,
                    workbook=workbook,
                    tenant_id=tenant_id,
                    bucket_id=bucket_id,
                )
            )
    return documents


def _load_excel_chart_xml_documents(
    path: Path,
    *,
    workbook: Any,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    try:
        with ZipFile(path) as archive:
            chart_paths = sorted(
                name
                for name in archive.namelist()
                if name.startswith("xl/charts/chart") and name.endswith(".xml")
            )
            for chart_index, chart_path in enumerate(chart_paths, start=1):
                xml_text = archive.read(chart_path).decode("utf-8", errors="replace")
                chart_summary = _summarize_excel_chart_xml(xml_text)
                workbook_context = _workbook_context(workbook)
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Chart XML: {chart_path}",
                        "Block type: excel_chart",
                        f"Chart type: {chart_summary['chart_type']}",
                        f"Chart labels: {', '.join(chart_summary['labels']) if chart_summary['labels'] else 'not available'}",
                        f"Chart references: {', '.join(chart_summary['references']) if chart_summary['references'] else 'not available'}",
                        "Visible workbook context:",
                        workbook_context,
                    ]
                )
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:chart_xml:{chart_index}",
                        title=f"{path.stem} - chart XML {chart_index}",
                        content=content,
                        source_type="excel_chart",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "block_type": "excel_chart",
                            "chart_index": chart_index,
                            "chart_type": chart_summary["chart_type"],
                            "chart_xml_path": chart_path,
                            "chart_labels": chart_summary["labels"],
                            "chart_references": chart_summary["references"],
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    except BadZipFile:
        return []
    return documents


def _worksheet_context(worksheet: Any, *, max_rows: int = 20) -> str:
    rows: list[str] = []
    for row in worksheet.iter_rows(values_only=True):
        values = [str(value) for value in row if value is not None and str(value).strip()]
        if not values:
            continue
        rows.append(" | ".join(values))
        if len(rows) >= max_rows:
            break
    return "\n".join(rows)


def _workbook_context(workbook: Any, *, max_rows_per_sheet: int = 12) -> str:
    sections: list[str] = []
    for worksheet in workbook.worksheets:
        context = _worksheet_context(worksheet, max_rows=max_rows_per_sheet)
        if context:
            sections.append(f"Sheet: {worksheet.title}\n{context}")
    return "\n\n".join(sections)


def _chart_references(chart: Any) -> list[str]:
    references: list[str] = []
    for series in getattr(chart, "series", []) or []:
        for ref_path in ("cat.numRef.f", "val.numRef.f", "xVal.numRef.f", "yVal.numRef.f"):
            value = _nested_attr(series, ref_path)
            if value and value not in references:
                references.append(str(value))
    return references


def _nested_attr(value: Any, path: str) -> Any:
    current = value
    for attribute in path.split("."):
        current = getattr(current, attribute, None)
        if current is None:
            return None
    return current


def _summarize_excel_chart_xml(xml_text: str) -> dict[str, Any]:
    root = ET.fromstring(xml_text)
    labels: list[str] = []
    references: list[str] = []
    chart_types: list[str] = []

    for element in root.iter():
        tag = _xml_local_name(element.tag)
        text = (element.text or "").strip()
        if tag in {"v", "f"} and text:
            target = references if "!" in text or text.startswith("_xlchart") else labels
            if text not in target:
                target.append(text)
        layout_id = element.attrib.get("layoutId")
        if layout_id and layout_id not in chart_types:
            chart_types.append(layout_id)
        if tag.endswith("Chart") and tag not in chart_types:
            chart_types.append(tag)

    return {
        "chart_type": ", ".join(chart_types[:5]) or "chartXml",
        "labels": labels[:20],
        "references": references[:20],
    }


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]

