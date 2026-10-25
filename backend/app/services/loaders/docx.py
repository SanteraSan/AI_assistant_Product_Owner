from __future__ import annotations

from pathlib import Path

from docx import Document

from app.services.loaders.common import (
    DOCX_PARAGRAPH_WINDOW_OVERLAP,
    DOCX_PARAGRAPH_WINDOW_SIZE,
    RawDocument,
    _domain_for_path,
    _features_from_text,
    _stable_document_id,
)


def _load_docx_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.docx")):
        docx_document = Document(path)
        paragraphs = [
            (paragraph_index, paragraph.text.strip())
            for paragraph_index, paragraph in enumerate(docx_document.paragraphs)
            if paragraph.text.strip()
        ]
        block_index = 0
        for position, (paragraph_index, content) in enumerate(paragraphs):
            block_index += 1
            previous_content = paragraphs[position - 1][1] if position > 0 else ""
            next_content = paragraphs[position + 1][1] if position + 1 < len(paragraphs) else ""
            documents.append(
                RawDocument(
                    id=f"{_stable_document_id(path)}:paragraph:{paragraph_index}",
                    title=f"{path.stem} - paragraph {paragraph_index + 1}",
                    content="\n".join(
                        [
                            f"File: {path.name}",
                            f"Block type: paragraph",
                            f"Previous paragraph: {previous_content}" if previous_content else "",
                            f"Current paragraph: {content}",
                            f"Next paragraph: {next_content}" if next_content else "",
                        ]
                    ),
                    source_type="docx",
                    source_path=str(path),
                    domain=_domain_for_path(path),
                    feature=_features_from_text(content),
                    metadata={
                        "file_name": path.name,
                        "block_type": "paragraph",
                        "block_index": block_index,
                        "paragraph_index": paragraph_index,
                    },
                    tenant_id=tenant_id,
                    bucket_id=bucket_id,
                )
            )

        documents.extend(
            _build_docx_paragraph_windows(
                path,
                paragraphs,
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )

        for table_index, table in enumerate(docx_document.tables):
            for row_index, row in enumerate(table.rows):
                row_values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if not row_values:
                    continue
                block_index += 1
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Block type: table_row",
                        f"Table: {table_index + 1}",
                        f"Row values: {' | '.join(row_values)}",
                    ]
                )
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:table:{table_index}:row:{row_index}",
                        title=f"{path.stem} - table {table_index + 1} row {row_index + 1}",
                        content=content,
                        source_type="docx",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "block_type": "table_row",
                            "block_index": block_index,
                            "table_index": table_index,
                            "table_row_index": row_index,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    return documents


def _build_docx_paragraph_windows(
    path: Path,
    paragraphs: list[tuple[int, str]],
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    if len(paragraphs) < 2:
        return []

    documents: list[RawDocument] = []
    step = max(1, DOCX_PARAGRAPH_WINDOW_SIZE - DOCX_PARAGRAPH_WINDOW_OVERLAP)
    for window_index, start in enumerate(range(0, len(paragraphs), step)):
        window = paragraphs[start : start + DOCX_PARAGRAPH_WINDOW_SIZE]
        if not window:
            continue

        paragraph_start_index = window[0][0]
        paragraph_end_index = window[-1][0]
        paragraph_lines = [
            f"Paragraph {paragraph_index + 1}: {content}"
            for paragraph_index, content in window
        ]
        content = "\n".join(
            [
                f"File: {path.name}",
                "Block type: paragraph_window",
                f"Paragraph range: {paragraph_start_index + 1}-{paragraph_end_index + 1}",
                *paragraph_lines,
            ]
        )
        documents.append(
            RawDocument(
                id=f"{_stable_document_id(path)}:paragraph_window:{window_index}",
                title=f"{path.stem} - paragraphs {paragraph_start_index + 1}-{paragraph_end_index + 1}",
                content=content,
                source_type="docx",
                source_path=str(path),
                domain=_domain_for_path(path),
                feature=_features_from_text(content),
                metadata={
                    "file_name": path.name,
                    "block_type": "paragraph_window",
                    "window_index": window_index,
                    "paragraph_start_index": paragraph_start_index,
                    "paragraph_end_index": paragraph_end_index,
                    "paragraph_count": len(window),
                },
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents

