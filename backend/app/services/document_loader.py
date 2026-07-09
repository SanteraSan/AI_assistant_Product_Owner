import json
import warnings
from dataclasses import dataclass, field
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from typing import Any, Iterator
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZipFile

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import pandas as pd
from PIL import Image
import pytesseract
from pypdf import PdfReader
import yaml

DEFAULT_TENANT_ID = "local_demo"
DEFAULT_BUCKET_ID = "taskflow_seed"
INDEXED_STATUS = "indexed"
INGESTION_MANIFEST_FILE_NAME = "ingestion_manifest.json"
DOCX_PARAGRAPH_WINDOW_SIZE = 8
DOCX_PARAGRAPH_WINDOW_OVERLAP = 2
IMAGE_OCR_EXTENSIONS = (".png", ".jpg", ".jpeg")
IMAGE_OCR_LANGUAGES = "rus+eng"
OFFICE_IMAGE_EXTENSIONS = (".png", ".jpg", ".jpeg")
DOCX_MEDIA_PREFIX = "word/media/"
XLSX_MEDIA_PREFIX = "xl/media/"
LEGACY_XLS_EMBEDDED_IMAGE_LIMIT = 50
SCANNED_TABLE_MIN_GRID_LINES = 3
SCANNED_TABLE_LINE_RATIO = 0.25
SCANNED_TABLE_CELL_PADDING = 4
OFFICE_XML_NAMESPACES = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


@dataclass(frozen=True)
class RawDocument:
    id: str
    title: str
    content: str
    source_type: str
    source_path: str
    domain: str
    feature: list[str] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    tenant_id: str = DEFAULT_TENANT_ID
    bucket_id: str = DEFAULT_BUCKET_ID
    processing_status: str = INDEXED_STATUS


@dataclass(frozen=True)
class EmbeddedImage:
    parent_path: Path
    embedded_path: str
    image_index: int
    content: bytes
    parent_source_type: str
    metadata: dict[str, Any] = field(default_factory=dict)


def load_raw_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str = DEFAULT_TENANT_ID,
    bucket_id: str = DEFAULT_BUCKET_ID,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    documents.extend(_load_markdown_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_text_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_json_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_pdf_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_docx_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_image_ocr_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_office_embedded_image_ocr_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_scanned_table_image_ocr_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_excel_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_excel_chart_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_csv_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents.extend(_load_openapi_documents(raw_data_dir, tenant_id=tenant_id, bucket_id=bucket_id))
    documents = _apply_ingestion_manifest(raw_data_dir, documents)
    return [document for document in documents if document.content.strip()]


def _load_markdown_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.md")):
        content = path.read_text(encoding="utf-8")
        documents.append(
            RawDocument(
                id=_stable_document_id(path),
                title=_title_from_markdown(content) or path.stem.replace("_", " ").title(),
                content=content,
                source_type=_source_type_for_path(path),
                source_path=str(path),
                domain=_domain_for_path(path),
                feature=_features_from_text(content),
                metadata={"file_name": path.name},
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _load_text_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.txt")):
        content = path.read_text(encoding="utf-8")
        documents.append(
            RawDocument(
                id=_stable_document_id(path),
                title=path.stem.replace("_", " ").title(),
                content=content,
                source_type=_source_type_for_path(path),
                source_path=str(path),
                domain=_domain_for_path(path),
                feature=_features_from_text(content),
                metadata={"file_name": path.name},
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _load_json_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.json")):
        if path.name == INGESTION_MANIFEST_FILE_NAME:
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        content = _json_to_text(data)
        documents.append(
            RawDocument(
                id=_stable_document_id(path),
                title=_title_from_json(data) or path.stem.replace("_", " ").title(),
                content=content,
                source_type=_source_type_for_path(path),
                source_path=str(path),
                domain=_domain_for_path(path),
                feature=_features_from_text(content),
                metadata={"file_name": path.name},
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _load_pdf_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.pdf")):
        reader = PdfReader(path)
        page_count = len(reader.pages)
        for page_index, page in enumerate(reader.pages, start=1):
            content = (page.extract_text() or "").strip()
            documents.append(
                RawDocument(
                    id=f"{_stable_document_id(path)}:page:{page_index}",
                    title=f"{path.stem.replace('_', ' ').title()} - page {page_index}",
                    content=content,
                    source_type="pdf",
                    source_path=str(path),
                    domain=_domain_for_path(path),
                    feature=_features_from_text(content),
                    metadata={
                        "file_name": path.name,
                        "page_number": page_index,
                        "page_count": page_count,
                    },
                    tenant_id=tenant_id,
                    bucket_id=bucket_id,
                )
            )
    return documents


def _load_csv_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.csv")):
        frame = pd.read_csv(path)
        for index, row in frame.fillna("").iterrows():
            row_data = {key: str(value) for key, value in row.to_dict().items()}
            title = row_data.get("subject") or row_data.get("review_id") or row_data.get("metric_type") or path.stem
            feature = _split_features(row_data.get("feature", ""))
            documents.append(
                RawDocument(
                    id=f"{_stable_document_id(path)}:{index}",
                    title=str(title),
                    content=_row_to_text(row_data),
                    source_type=_source_type_for_path(path),
                    source_path=str(path),
                    domain=_domain_for_path(path),
                    feature=feature,
                    metadata={"row_index": int(index), "file_name": path.name, **row_data},
                    tenant_id=tenant_id,
                    bucket_id=bucket_id,
                )
            )
    return documents


def _load_docx_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.docx")):
        docx_document = Document(path)
        paragraphs = [
            (paragraph_index, paragraph.text.strip())
            for paragraph_index, paragraph in enumerate(docx_document.paragraphs)
            if paragraph.text.strip()
        ]
        block_index = 0
        for position, (paragraph_index, content) in enumerate(paragraphs):
            block_index += 1
            previous_content = paragraphs[position - 1][1] if position > 0 else ""
            next_content = paragraphs[position + 1][1] if position + 1 < len(paragraphs) else ""
            documents.append(
                RawDocument(
                    id=f"{_stable_document_id(path)}:paragraph:{paragraph_index}",
                    title=f"{path.stem} - paragraph {paragraph_index + 1}",
                    content="\n".join(
                        [
                            f"File: {path.name}",
                            f"Block type: paragraph",
                            f"Previous paragraph: {previous_content}" if previous_content else "",
                            f"Current paragraph: {content}",
                            f"Next paragraph: {next_content}" if next_content else "",
                        ]
                    ),
                    source_type="docx",
                    source_path=str(path),
                    domain=_domain_for_path(path),
                    feature=_features_from_text(content),
                    metadata={
                        "file_name": path.name,
                        "block_type": "paragraph",
                        "block_index": block_index,
                        "paragraph_index": paragraph_index,
                    },
                    tenant_id=tenant_id,
                    bucket_id=bucket_id,
                )
            )

        documents.extend(
            _build_docx_paragraph_windows(
                path,
                paragraphs,
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )

        for table_index, table in enumerate(docx_document.tables):
            for row_index, row in enumerate(table.rows):
                row_values = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if not row_values:
                    continue
                block_index += 1
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Block type: table_row",
                        f"Table: {table_index + 1}",
                        f"Row values: {' | '.join(row_values)}",
                    ]
                )
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:table:{table_index}:row:{row_index}",
                        title=f"{path.stem} - table {table_index + 1} row {row_index + 1}",
                        content=content,
                        source_type="docx",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "block_type": "table_row",
                            "block_index": block_index,
                            "table_index": table_index,
                            "table_row_index": row_index,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    return documents


def _load_image_ocr_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    image_paths = _image_paths(raw_data_dir)
    if not image_paths:
        return documents

    if not _tesseract_available():
        warnings.warn(
            (
                "Tesseract OCR is not available. Image OCR ingestion skipped. "
                "Install system packages: tesseract-ocr tesseract-ocr-rus tesseract-ocr-eng."
            ),
            RuntimeWarning,
            stacklevel=2,
        )
        return documents

    for path in image_paths:
        with Image.open(path) as image:
            width, height = image.size
            image_format = image.format or path.suffix.lstrip(".").upper()
            content = pytesseract.image_to_string(image, lang=IMAGE_OCR_LANGUAGES).strip()
        if not content:
            continue

        documents.append(
            RawDocument(
                id=f"{_stable_document_id(path)}:ocr",
                title=f"{path.stem} - OCR text",
                content="\n".join(
                    [
                        f"File: {path.name}",
                        "Block type: image_ocr",
                        "OCR text:",
                        content,
                    ]
                ),
                source_type="image_ocr",
                source_path=str(path),
                domain=_domain_for_path(path),
                feature=_features_from_text(content),
                metadata={
                    "file_name": path.name,
                    "block_type": "image_ocr",
                    "image_width": width,
                    "image_height": height,
                    "image_format": image_format,
                    "ocr_engine": "tesseract",
                    "ocr_languages": IMAGE_OCR_LANGUAGES,
                },
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _image_paths(raw_data_dir: Path) -> list[Path]:
    return sorted(
        path
        for extension in IMAGE_OCR_EXTENSIONS
        for path in raw_data_dir.rglob(f"*{extension}")
    )


def _load_office_embedded_image_ocr_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    if not _tesseract_available():
        return []

    documents: list[RawDocument] = []
    for embedded_image in iter_office_embedded_images(raw_data_dir):
        try:
            image, width, height, image_format = _open_embedded_image(embedded_image)
            content = pytesseract.image_to_string(image, lang=IMAGE_OCR_LANGUAGES).strip()
        except Exception as exc:
            warnings.warn(
                f"Embedded image OCR skipped for {embedded_image.parent_path}:{embedded_image.embedded_path}: {exc}",
                RuntimeWarning,
                stacklevel=2,
            )
            continue
        if not content:
            continue

        parent_path = embedded_image.parent_path
        documents.append(
            RawDocument(
                id=f"{_stable_document_id(parent_path)}:embedded_image:{embedded_image.image_index}:ocr",
                title=f"{parent_path.stem} - embedded image {embedded_image.image_index} OCR text",
                content="\n".join(
                    [
                        f"File: {parent_path.name}",
                        f"Embedded image: {embedded_image.embedded_path}",
                        "Block type: image_ocr",
                        *_embedded_image_context_lines(embedded_image.metadata),
                        "OCR text:",
                        content,
                    ]
                ),
                source_type="image_ocr",
                source_path=str(parent_path),
                domain=_domain_for_path(parent_path),
                feature=_features_from_text(content),
                metadata={
                    "file_name": parent_path.name,
                    "block_type": "image_ocr",
                    "parent_source_type": embedded_image.parent_source_type,
                    "embedded_path": embedded_image.embedded_path,
                    "embedded_image_index": embedded_image.image_index,
                    "image_width": width,
                    "image_height": height,
                    "image_format": image_format,
                    "ocr_engine": "tesseract",
                    "ocr_languages": IMAGE_OCR_LANGUAGES,
                    **embedded_image.metadata,
                },
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _load_scanned_table_image_ocr_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    if not _tesseract_available():
        return []

    documents: list[RawDocument] = []
    for embedded_image in iter_scanned_table_images(raw_data_dir):
        try:
            image, width, height, image_format = _open_embedded_image(embedded_image)
            content = pytesseract.image_to_string(image, lang=IMAGE_OCR_LANGUAGES).strip()
        except Exception as exc:
            warnings.warn(
                f"Scanned table image OCR skipped for {embedded_image.parent_path}:{embedded_image.embedded_path}: {exc}",
                RuntimeWarning,
                stacklevel=2,
            )
            continue
        if not content:
            continue

        parent_path = embedded_image.parent_path
        documents.append(
            RawDocument(
                id=f"{_stable_document_id(parent_path)}:scanned_table_image:{embedded_image.image_index}:ocr",
                title=f"{parent_path.stem} - scanned table image {embedded_image.image_index} OCR text",
                content="\n".join(
                    [
                        f"File: {parent_path.name}",
                        f"Embedded image: {embedded_image.embedded_path}",
                        "Block type: image_ocr",
                        *_embedded_image_context_lines(embedded_image.metadata),
                        "OCR text:",
                        content,
                    ]
                ),
                source_type="image_ocr",
                source_path=str(parent_path),
                domain=_domain_for_path(parent_path),
                feature=_features_from_text(content),
                metadata={
                    "file_name": parent_path.name,
                    "block_type": "image_ocr",
                    "parent_source_type": embedded_image.parent_source_type,
                    "embedded_path": embedded_image.embedded_path,
                    "embedded_image_index": embedded_image.image_index,
                    "image_width": width,
                    "image_height": height,
                    "image_format": image_format,
                    "ocr_engine": "tesseract",
                    "ocr_languages": IMAGE_OCR_LANGUAGES,
                    **embedded_image.metadata,
                },
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _tesseract_available() -> bool:
    try:
        pytesseract.get_tesseract_version()
    except pytesseract.TesseractNotFoundError:
        return False
    return True


def iter_office_embedded_images(raw_data_dir: Path) -> list[EmbeddedImage]:
    images: list[EmbeddedImage] = []
    for path in sorted(
        list(raw_data_dir.rglob("*.docx"))
        + list(raw_data_dir.rglob("*.xlsx"))
        + list(raw_data_dir.rglob("*.xls"))
    ):
        images.extend(_embedded_images_from_office_file(path))
    return images


def iter_scanned_table_images(raw_data_dir: Path) -> list[EmbeddedImage]:
    images: list[EmbeddedImage] = []
    for path in _image_paths(raw_data_dir):
        images.extend(_scanned_table_images_from_image_file(path))
    return images


def _embedded_images_from_office_file(path: Path) -> list[EmbeddedImage]:
    if path.suffix.lower() == ".xls":
        return _embedded_images_from_legacy_xls(path)
    if path.suffix.lower() == ".xlsx":
        return _embedded_images_from_xlsx(path)
    if path.suffix.lower() == ".docx":
        return _embedded_images_from_docx(path)

    return []


def _embedded_images_from_docx(path: Path) -> list[EmbeddedImage]:
    try:
        with ZipFile(path) as archive:
            metadata_by_path = _docx_image_metadata_by_embedded_path(archive)
            media_paths = sorted(
                name
                for name in archive.namelist()
                if name.startswith(DOCX_MEDIA_PREFIX)
                and Path(name).suffix.lower() in OFFICE_IMAGE_EXTENSIONS
            )
            return [
                EmbeddedImage(
                    parent_path=path,
                    embedded_path=media_path,
                    image_index=image_index,
                    content=archive.read(media_path),
                    parent_source_type="docx",
                    metadata=metadata_by_path.get(media_path, {}),
                )
                for image_index, media_path in enumerate(media_paths, start=1)
            ]
    except BadZipFile:
        warnings.warn(
            f"Office embedded images skipped for non-zip Office file: {path}",
            RuntimeWarning,
            stacklevel=2,
        )
        return []


def _embedded_images_from_xlsx(path: Path) -> list[EmbeddedImage]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    images: list[EmbeddedImage] = []
    image_index = 0
    for worksheet in workbook.worksheets:
        for image in getattr(worksheet, "_images", []):
            image_index += 1
            content = image._data()
            extension = _image_extension(getattr(image, "format", None), content)
            images.append(
                EmbeddedImage(
                    parent_path=path,
                    embedded_path=f"{XLSX_MEDIA_PREFIX}anchored_image{image_index}.{extension}",
                    image_index=image_index,
                    content=content,
                    parent_source_type="xlsx",
                    metadata=_xlsx_image_anchor_metadata(worksheet, image),
                )
            )
    return images


def _embedded_images_from_legacy_xls(path: Path) -> list[EmbeddedImage]:
    data = path.read_bytes()
    image_blobs = [
        *list(_iter_jpeg_blobs(data)),
        *list(_iter_png_blobs(data)),
    ]
    images: list[EmbeddedImage] = []
    for image_index, (extension, content) in enumerate(
        image_blobs[:LEGACY_XLS_EMBEDDED_IMAGE_LIMIT],
        start=1,
    ):
        images.append(
            EmbeddedImage(
                parent_path=path,
                embedded_path=f"legacy-binary/image{image_index}.{extension}",
                image_index=image_index,
                content=content,
                parent_source_type="xls",
            )
        )
    return images


def _scanned_table_images_from_image_file(path: Path) -> list[EmbeddedImage]:
    try:
        with Image.open(path) as image:
            rgb_image = image.convert("RGB")
            vertical_lines, horizontal_lines = _detect_table_grid_lines(rgb_image)
            if (
                len(vertical_lines) < SCANNED_TABLE_MIN_GRID_LINES
                or len(horizontal_lines) < SCANNED_TABLE_MIN_GRID_LINES
            ):
                return []

            image_column_index = len(vertical_lines) - 2
            embedded_images: list[EmbeddedImage] = []
            for row_position, row_index in enumerate(range(1, len(horizontal_lines) - 1), start=1):
                row_box = _cell_box(
                    vertical_lines[0],
                    horizontal_lines[row_index],
                    vertical_lines[-1],
                    horizontal_lines[row_index + 1],
                )
                image_cell_box = _cell_box(
                    vertical_lines[image_column_index],
                    horizontal_lines[row_index],
                    vertical_lines[image_column_index + 1],
                    horizontal_lines[row_index + 1],
                )
                if not _cell_has_visual_content(rgb_image, image_cell_box):
                    continue

                crop = rgb_image.crop(image_cell_box)
                content = _encode_image_bytes(crop)
                linked_text = _ocr_image_region(rgb_image.crop(row_box))
                if not linked_text.strip():
                    linked_text = _ocr_image_region(
                        rgb_image.crop(
                            _cell_box(
                                vertical_lines[0],
                                horizontal_lines[row_index],
                                vertical_lines[image_column_index],
                                horizontal_lines[row_index + 1],
                            )
                        )
                    )

                embedded_images.append(
                    EmbeddedImage(
                        parent_path=path,
                        embedded_path=f"scanned-table/table1-row{row_position}-image.png",
                        image_index=row_position,
                        content=content,
                        parent_source_type="scanned_table",
                        metadata={
                            "anchor_type": "scanned_table_cell",
                            "table_index": 0,
                            "table_row_index": row_position,
                            "table_cell_index": image_column_index,
                            "image_cell_bbox": _bbox_metadata(image_cell_box),
                            "table_row_bbox": _bbox_metadata(row_box),
                            "linked_text": linked_text,
                        },
                    )
                )
            return embedded_images
    except Exception as exc:
        warnings.warn(
            f"Scanned table image extraction skipped for {path}: {exc}",
            RuntimeWarning,
            stacklevel=2,
        )
        return []


def _detect_table_grid_lines(image: Image.Image) -> tuple[list[int], list[int]]:
    width, height = image.size
    pixels = image.load()

    horizontal_candidates = [
        y
        for y in range(height)
        if _line_pixel_longest_run(pixels, width, y, horizontal=True) / max(1, width)
        >= SCANNED_TABLE_LINE_RATIO
    ]
    vertical_candidates = [
        x
        for x in range(width)
        if _line_pixel_longest_run(pixels, height, x, horizontal=False) / max(1, height)
        >= SCANNED_TABLE_LINE_RATIO
    ]
    return (
        _line_group_centers(vertical_candidates),
        _line_group_centers(horizontal_candidates),
    )


def _line_pixel_longest_run(pixels: Any, length: int, fixed_position: int, *, horizontal: bool) -> int:
    longest_run = 0
    current_run = 0
    for position in range(length):
        pixel = pixels[position, fixed_position] if horizontal else pixels[fixed_position, position]
        if _is_table_line_pixel(pixel):
            current_run += 1
            longest_run = max(longest_run, current_run)
        else:
            current_run = 0
    return longest_run


def _is_table_line_pixel(pixel: tuple[int, int, int]) -> bool:
    red, green, blue = pixel
    is_blue_grid = blue >= 120 and red <= 140 and green <= 170
    is_dark_grid = red <= 80 and green <= 80 and blue <= 80
    return is_blue_grid or is_dark_grid


def _line_group_centers(positions: list[int]) -> list[int]:
    if not positions:
        return []

    groups: list[list[int]] = [[positions[0]]]
    for position in positions[1:]:
        if position - groups[-1][-1] <= 2:
            groups[-1].append(position)
        else:
            groups.append([position])
    return [group[len(group) // 2] for group in groups]


def _cell_box(left: int, top: int, right: int, bottom: int) -> tuple[int, int, int, int]:
    return (
        left + SCANNED_TABLE_CELL_PADDING,
        top + SCANNED_TABLE_CELL_PADDING,
        max(left + SCANNED_TABLE_CELL_PADDING + 1, right - SCANNED_TABLE_CELL_PADDING),
        max(top + SCANNED_TABLE_CELL_PADDING + 1, bottom - SCANNED_TABLE_CELL_PADDING),
    )


def _cell_has_visual_content(image: Image.Image, box: tuple[int, int, int, int]) -> bool:
    crop = image.crop(box)
    width, height = crop.size
    if width <= 0 or height <= 0:
        return False

    rgb_crop = crop.convert("RGB")
    pixels = rgb_crop.load()
    non_background = 0
    for y in range(height):
        for x in range(width):
            red, green, blue = pixels[x, y]
            if _is_foreground_pixel(red, green, blue):
                non_background += 1
    return non_background / max(1, width * height) >= 0.08


def _is_foreground_pixel(red: int, green: int, blue: int) -> bool:
    is_white = red >= 245 and green >= 245 and blue >= 245
    is_light_blue_fill = red >= 190 and green >= 205 and blue >= 220
    is_grid = _is_table_line_pixel((red, green, blue))
    return not (is_white or is_light_blue_fill or is_grid)


def _ocr_image_region(image: Image.Image) -> str:
    if not _tesseract_available():
        return ""
    return pytesseract.image_to_string(image, lang=IMAGE_OCR_LANGUAGES).strip()


def _encode_image_bytes(image: Image.Image) -> bytes:
    buffer = BytesIO()
    image.convert("RGB").save(buffer, format="PNG")
    return buffer.getvalue()


def _bbox_metadata(box: tuple[int, int, int, int]) -> dict[str, int]:
    left, top, right, bottom = box
    return {
        "left": left,
        "top": top,
        "right": right,
        "bottom": bottom,
    }


def _docx_image_metadata_by_embedded_path(archive: ZipFile) -> dict[str, dict[str, Any]]:
    try:
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        relationships = _docx_relationship_targets(archive)
    except Exception:
        return {}

    metadata_by_path: dict[str, dict[str, Any]] = {}
    body = document_xml.find("w:body", OFFICE_XML_NAMESPACES)
    if body is None:
        return metadata_by_path

    paragraph_index = -1
    table_index = -1
    previous_paragraph_text = ""
    for child in list(body):
        if child.tag == f"{{{OFFICE_XML_NAMESPACES['w']}}}p":
            paragraph_index += 1
            paragraph_text = _xml_text(child)
            linked_text = paragraph_text or previous_paragraph_text
            for relationship_id in _embedded_relationship_ids(child):
                embedded_path = _docx_embedded_path(relationships.get(relationship_id, ""))
                if embedded_path:
                    metadata_by_path.setdefault(
                        embedded_path,
                        {
                            "anchor_type": "docx_paragraph",
                            "paragraph_index": paragraph_index,
                            "paragraph_text": paragraph_text,
                            "previous_paragraph_text": previous_paragraph_text,
                            "linked_text": linked_text,
                        },
                    )
            if paragraph_text:
                previous_paragraph_text = paragraph_text
            continue

        if child.tag != f"{{{OFFICE_XML_NAMESPACES['w']}}}tbl":
            continue

        table_index += 1
        for row_index, row in enumerate(child.findall("w:tr", OFFICE_XML_NAMESPACES)):
            cells = row.findall("w:tc", OFFICE_XML_NAMESPACES)
            row_values = [_xml_text(cell) for cell in cells]
            row_text = " | ".join(value for value in row_values if value)
            for cell_index, cell in enumerate(cells):
                cell_text = row_values[cell_index] if cell_index < len(row_values) else ""
                for relationship_id in _embedded_relationship_ids(cell):
                    embedded_path = _docx_embedded_path(relationships.get(relationship_id, ""))
                    if embedded_path:
                        metadata_by_path.setdefault(
                            embedded_path,
                            {
                                "anchor_type": "docx_table_cell",
                                "table_index": table_index,
                                "table_row_index": row_index,
                                "table_cell_index": cell_index,
                                "table_row_text": row_text,
                                "table_cell_text": cell_text,
                                "linked_text": row_text or cell_text,
                            },
                        )
    return metadata_by_path


def _docx_relationship_targets(archive: ZipFile) -> dict[str, str]:
    relationships_xml = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    return {
        relationship.attrib["Id"]: relationship.attrib.get("Target", "")
        for relationship in relationships_xml.findall(
            "rel:Relationship",
            OFFICE_XML_NAMESPACES,
        )
        if relationship.attrib.get("Id")
    }


def _embedded_relationship_ids(node: ET.Element) -> list[str]:
    relationship_ids: list[str] = []
    for blip in node.findall(".//a:blip", OFFICE_XML_NAMESPACES):
        relationship_id = blip.attrib.get(f"{{{OFFICE_XML_NAMESPACES['r']}}}embed")
        if relationship_id:
            relationship_ids.append(relationship_id)
    return relationship_ids


def _docx_embedded_path(target: str) -> str:
    if not target:
        return ""
    if target.startswith("word/"):
        return target
    return f"word/{target.lstrip('/')}"


def _xml_text(node: ET.Element) -> str:
    return "".join(
        text.text or ""
        for text in node.findall(".//w:t", OFFICE_XML_NAMESPACES)
    ).strip()


def _xlsx_image_anchor_metadata(worksheet: Any, image: Any) -> dict[str, Any]:
    marker = getattr(getattr(image, "anchor", None), "_from", None)
    if marker is None:
        return {
            "anchor_type": "xlsx_unknown",
            "sheet_name": worksheet.title,
        }

    row_number = int(marker.row) + 1
    column_number = int(marker.col) + 1
    row_text = _worksheet_row_text_by_number(worksheet, row_number)
    cell_value = worksheet.cell(row=row_number, column=column_number).value
    return {
        "anchor_type": "xlsx_cell",
        "sheet_name": worksheet.title,
        "anchor_row": row_number,
        "anchor_col": column_number,
        "anchor_cell": f"{get_column_letter(column_number)}{row_number}",
        "anchor_cell_value": str(cell_value).strip() if cell_value is not None else "",
        "nearby_row_text": row_text,
        "linked_text": row_text,
    }


def _worksheet_row_text_by_number(worksheet: Any, row_number: int) -> str:
    values = [
        str(cell.value).strip()
        for cell in worksheet[row_number]
        if cell.value not in (None, "")
    ]
    return " | ".join(values)


def _image_extension(image_format: str | None, content: bytes) -> str:
    normalized_format = (image_format or "").lower()
    if normalized_format in {"jpeg", "jpg"}:
        return "jpg"
    if normalized_format == "png":
        return "png"
    if content.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    return normalized_format or "bin"


def _embedded_image_context_lines(metadata: dict[str, Any]) -> list[str]:
    lines: list[str] = []
    if metadata.get("anchor_type"):
        lines.append(f"Anchor type: {metadata['anchor_type']}")
    if metadata.get("sheet_name"):
        lines.append(f"Sheet: {metadata['sheet_name']}")
    if metadata.get("anchor_cell"):
        lines.append(f"Anchor cell: {metadata['anchor_cell']}")
    if metadata.get("table_index") is not None:
        lines.append(
            "Table position: "
            f"table {int(metadata['table_index']) + 1}, "
            f"row {int(metadata.get('table_row_index', 0)) + 1}, "
            f"cell {int(metadata.get('table_cell_index', 0)) + 1}"
        )
    linked_text = str(metadata.get("linked_text") or "").strip()
    if linked_text:
        lines.append(f"Linked text: {linked_text}")
    return lines


def _iter_jpeg_blobs(data: bytes) -> Iterator[tuple[str, bytes]]:
    position = 0
    while True:
        start = data.find(b"\xff\xd8\xff", position)
        if start < 0:
            return
        end = data.find(b"\xff\xd9", start)
        if end < 0:
            return
        blob = data[start : end + 2]
        if _valid_image_blob(blob):
            yield "jpg", blob
        position = end + 2


def _iter_png_blobs(data: bytes) -> Iterator[tuple[str, bytes]]:
    signature = b"\x89PNG\r\n\x1a\n"
    position = 0
    while True:
        start = data.find(signature, position)
        if start < 0:
            return
        end = data.find(b"IEND", start)
        if end < 0:
            return
        blob = data[start : end + 8]
        if _valid_image_blob(blob):
            yield "png", blob
        position = end + 8


def _valid_image_blob(blob: bytes) -> bool:
    try:
        with Image.open(BytesIO(blob)) as image:
            image.load()
    except Exception:
        return False
    return True


def _open_embedded_image(embedded_image: EmbeddedImage) -> tuple[Image.Image, int, int, str]:
    image = Image.open(BytesIO(embedded_image.content))
    width, height = image.size
    image_format = image.format or Path(embedded_image.embedded_path).suffix.lstrip(".").upper()
    return image, width, height, image_format


def _build_docx_paragraph_windows(
    path: Path,
    paragraphs: list[tuple[int, str]],
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    if len(paragraphs) < 2:
        return []

    documents: list[RawDocument] = []
    step = max(1, DOCX_PARAGRAPH_WINDOW_SIZE - DOCX_PARAGRAPH_WINDOW_OVERLAP)
    for window_index, start in enumerate(range(0, len(paragraphs), step)):
        window = paragraphs[start : start + DOCX_PARAGRAPH_WINDOW_SIZE]
        if not window:
            continue

        paragraph_start_index = window[0][0]
        paragraph_end_index = window[-1][0]
        paragraph_lines = [
            f"Paragraph {paragraph_index + 1}: {content}"
            for paragraph_index, content in window
        ]
        content = "\n".join(
            [
                f"File: {path.name}",
                "Block type: paragraph_window",
                f"Paragraph range: {paragraph_start_index + 1}-{paragraph_end_index + 1}",
                *paragraph_lines,
            ]
        )
        documents.append(
            RawDocument(
                id=f"{_stable_document_id(path)}:paragraph_window:{window_index}",
                title=f"{path.stem} - paragraphs {paragraph_start_index + 1}-{paragraph_end_index + 1}",
                content=content,
                source_type="docx",
                source_path=str(path),
                domain=_domain_for_path(path),
                feature=_features_from_text(content),
                metadata={
                    "file_name": path.name,
                    "block_type": "paragraph_window",
                    "window_index": window_index,
                    "paragraph_start_index": paragraph_start_index,
                    "paragraph_end_index": paragraph_end_index,
                    "paragraph_count": len(window),
                },
                tenant_id=tenant_id,
                bucket_id=bucket_id,
            )
        )
    return documents


def _load_excel_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    excel_paths = sorted(
        list(raw_data_dir.rglob("*.xlsx"))
        + list(raw_data_dir.rglob("*.xls"))
    )
    for path in excel_paths:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str)
        for sheet_name, frame in sheets.items():
            for index, row in frame.fillna("").iterrows():
                row_data = {key: str(value) for key, value in row.to_dict().items()}
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Sheet: {sheet_name}",
                        f"Row values: {_compact_row_values(row_data)}",
                        _row_to_text(row_data),
                    ]
                )
                title = f"{path.stem} - {sheet_name} - row {int(index) + 2}"
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:{sheet_name}:{index}",
                        title=title,
                        content=content,
                        source_type="excel_row",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "sheet_name": sheet_name,
                            "row_index": int(index),
                            "excel_row_number": int(index) + 2,
                            **row_data,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    return documents


def _load_excel_chart_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(raw_data_dir.rglob("*.xlsx")):
        workbook = load_workbook(path, data_only=True, read_only=False)
        path_documents: list[RawDocument] = []
        for worksheet in workbook.worksheets:
            charts = getattr(worksheet, "_charts", [])
            if not charts:
                continue
            sheet_context = _worksheet_context(worksheet)
            for chart_index, chart in enumerate(charts, start=1):
                chart_type = type(chart).__name__
                chart_refs = _chart_references(chart)
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Sheet: {worksheet.title}",
                        "Block type: excel_chart",
                        f"Chart type: {chart_type}",
                        f"Chart references: {', '.join(chart_refs) if chart_refs else 'not available'}",
                        "Visible sheet context:",
                        sheet_context,
                    ]
                )
                path_documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:{worksheet.title}:chart:{chart_index}",
                        title=f"{path.stem} - {worksheet.title} - chart {chart_index}",
                        content=content,
                        source_type="excel_chart",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "sheet_name": worksheet.title,
                            "block_type": "excel_chart",
                            "chart_index": chart_index,
                            "chart_type": chart_type,
                            "chart_references": chart_refs,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
        documents.extend(path_documents)
        if not path_documents:
            documents.extend(
                _load_excel_chart_xml_documents(
                    path,
                    workbook=workbook,
                    tenant_id=tenant_id,
                    bucket_id=bucket_id,
                )
            )
    return documents


def _load_excel_chart_xml_documents(
    path: Path,
    *,
    workbook: Any,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    try:
        with ZipFile(path) as archive:
            chart_paths = sorted(
                name
                for name in archive.namelist()
                if name.startswith("xl/charts/chart") and name.endswith(".xml")
            )
            for chart_index, chart_path in enumerate(chart_paths, start=1):
                xml_text = archive.read(chart_path).decode("utf-8", errors="replace")
                chart_summary = _summarize_excel_chart_xml(xml_text)
                workbook_context = _workbook_context(workbook)
                content = "\n".join(
                    [
                        f"File: {path.name}",
                        f"Chart XML: {chart_path}",
                        "Block type: excel_chart",
                        f"Chart type: {chart_summary['chart_type']}",
                        f"Chart labels: {', '.join(chart_summary['labels']) if chart_summary['labels'] else 'not available'}",
                        f"Chart references: {', '.join(chart_summary['references']) if chart_summary['references'] else 'not available'}",
                        "Visible workbook context:",
                        workbook_context,
                    ]
                )
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:chart_xml:{chart_index}",
                        title=f"{path.stem} - chart XML {chart_index}",
                        content=content,
                        source_type="excel_chart",
                        source_path=str(path),
                        domain=_domain_for_path(path),
                        feature=_features_from_text(content),
                        metadata={
                            "file_name": path.name,
                            "block_type": "excel_chart",
                            "chart_index": chart_index,
                            "chart_type": chart_summary["chart_type"],
                            "chart_xml_path": chart_path,
                            "chart_labels": chart_summary["labels"],
                            "chart_references": chart_summary["references"],
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    except BadZipFile:
        return []
    return documents


def _load_openapi_documents(
    raw_data_dir: Path,
    *,
    tenant_id: str,
    bucket_id: str,
) -> list[RawDocument]:
    documents: list[RawDocument] = []
    for path in sorted(list(raw_data_dir.rglob("*.yaml")) + list(raw_data_dir.rglob("*.yml"))):
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        paths = data.get("paths", {})
        for route, methods in paths.items():
            if not isinstance(methods, dict):
                continue
            for method, operation in methods.items():
                if not isinstance(operation, dict):
                    continue
                operation_id = operation.get("operationId", f"{method}_{route}")
                summary = operation.get("summary", "")
                tags = operation.get("tags", [])
                responses = ", ".join(str(code) for code in operation.get("responses", {}).keys())
                content = "\n".join(
                    [
                        f"Endpoint: {method.upper()} {route}",
                        f"Operation ID: {operation_id}",
                        f"Summary: {summary}",
                        f"Tags: {', '.join(tags) if isinstance(tags, list) else tags}",
                        f"Responses: {responses}",
                    ]
                )
                documents.append(
                    RawDocument(
                        id=f"{_stable_document_id(path)}:{method}:{route}",
                        title=f"{method.upper()} {route}",
                        content=content,
                        source_type="openapi",
                        source_path=str(path),
                        domain="tech_knowledge",
                        feature=_features_from_text(content),
                        metadata={
                            "method": method.upper(),
                            "path": route,
                            "operation_id": operation_id,
                            "tags": tags,
                            "file_name": path.name,
                        },
                        tenant_id=tenant_id,
                        bucket_id=bucket_id,
                    )
                )
    return documents


def _title_from_markdown(content: str) -> str | None:
    for line in content.splitlines():
        if line.startswith("# "):
            return line[2:].strip()
    return None


def _title_from_json(data: Any) -> str | None:
    if not isinstance(data, dict):
        return None
    for key in ("title", "name", "id"):
        value = data.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def _source_type_for_path(path: Path) -> str:
    parts = set(path.parts)
    if path.suffix == ".csv":
        if "business_metrics" in parts:
            return "metric_row"
        if "user_feedback" in parts:
            return "support_ticket"
    if "release_notes" in parts:
        return "incident_note" if "incident" in path.name else "release_note"
    if "tech_knowledge" in parts:
        return "markdown"
    return path.suffix.lstrip(".") or "text"


def _domain_for_path(path: Path) -> str:
    for domain in ("tech_knowledge", "user_feedback", "business_metrics", "release_notes"):
        if domain in path.parts:
            return domain
    return "documents"


def _stable_document_id(path: Path) -> str:
    return path.as_posix().replace("/", "_").replace(".", "_")


def _apply_ingestion_manifest(
    raw_data_dir: Path,
    documents: list[RawDocument],
) -> list[RawDocument]:
    manifest_path = raw_data_dir / INGESTION_MANIFEST_FILE_NAME
    if not manifest_path.exists():
        return documents

    data = json.loads(manifest_path.read_text(encoding="utf-8"))
    entries = data.get("documents", []) if isinstance(data, dict) else []
    overrides_by_path = {
        str(entry.get("path")): entry
        for entry in entries
        if isinstance(entry, dict) and entry.get("path")
    }
    if not overrides_by_path:
        return documents

    updated_documents: list[RawDocument] = []
    for document in documents:
        relative_path = Path(document.source_path).relative_to(raw_data_dir).as_posix()
        override = overrides_by_path.get(relative_path)
        if not override:
            updated_documents.append(document)
            continue

        metadata = dict(document.metadata)
        metadata.update(_as_dict(override.get("metadata")))
        updated_documents.append(
            replace(
                document,
                tenant_id=str(override.get("tenant_id") or document.tenant_id),
                bucket_id=str(override.get("bucket_id") or document.bucket_id),
                source_type=str(override.get("source_type") or document.source_type),
                feature=_as_str_list(override.get("features")) or document.feature,
                metadata=metadata,
            )
        )
    return updated_documents


def _features_from_text(text: str) -> list[str]:
    known_features = [
        "tasks",
        "projects",
        "sprints",
        "notifications",
        "permissions",
        "csv_import",
        "search",
        "reports",
        "webhooks",
        "integrations",
    ]
    lowered = text.lower()
    return [feature for feature in known_features if feature.lower() in lowered]


def _split_features(value: str) -> list[str]:
    return [item.strip() for item in value.split(",") if item.strip()]


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _as_str_list(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item) for item in value if str(item).strip()]
    if isinstance(value, str) and value.strip():
        return [value.strip()]
    return []


def _row_to_text(row: dict[str, str]) -> str:
    return "\n".join(f"{key}: {value}" for key, value in row.items() if value)


def _compact_row_values(row: dict[str, str]) -> str:
    return " | ".join(value.strip() for value in row.values() if value.strip())


def _worksheet_context(worksheet: Any, *, max_rows: int = 20) -> str:
    rows: list[str] = []
    for row in worksheet.iter_rows(values_only=True):
        values = [str(value) for value in row if value is not None and str(value).strip()]
        if not values:
            continue
        rows.append(" | ".join(values))
        if len(rows) >= max_rows:
            break
    return "\n".join(rows)


def _workbook_context(workbook: Any, *, max_rows_per_sheet: int = 12) -> str:
    sections: list[str] = []
    for worksheet in workbook.worksheets:
        context = _worksheet_context(worksheet, max_rows=max_rows_per_sheet)
        if context:
            sections.append(f"Sheet: {worksheet.title}\n{context}")
    return "\n\n".join(sections)


def _chart_references(chart: Any) -> list[str]:
    references: list[str] = []
    for series in getattr(chart, "series", []) or []:
        for ref_path in ("cat.numRef.f", "val.numRef.f", "xVal.numRef.f", "yVal.numRef.f"):
            value = _nested_attr(series, ref_path)
            if value and value not in references:
                references.append(str(value))
    return references


def _nested_attr(value: Any, path: str) -> Any:
    current = value
    for attribute in path.split("."):
        current = getattr(current, attribute, None)
        if current is None:
            return None
    return current


def _summarize_excel_chart_xml(xml_text: str) -> dict[str, Any]:
    root = ET.fromstring(xml_text)
    labels: list[str] = []
    references: list[str] = []
    chart_types: list[str] = []

    for element in root.iter():
        tag = _xml_local_name(element.tag)
        text = (element.text or "").strip()
        if tag in {"v", "f"} and text:
            target = references if "!" in text or text.startswith("_xlchart") else labels
            if text not in target:
                target.append(text)
        layout_id = element.attrib.get("layoutId")
        if layout_id and layout_id not in chart_types:
            chart_types.append(layout_id)
        if tag.endswith("Chart") and tag not in chart_types:
            chart_types.append(tag)

    return {
        "chart_type": ", ".join(chart_types[:5]) or "chartXml",
        "labels": labels[:20],
        "references": references[:20],
    }


def _xml_local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _json_to_text(data: Any) -> str:
    return json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True)
